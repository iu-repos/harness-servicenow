from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PRIMARY = RGBColor(13, 42, 75)
ACCENT = RGBColor(26, 163, 150)
TEXT_DARK = RGBColor(38, 50, 56)
TEXT_MID = RGBColor(77, 91, 102)
LIGHT_FILL = RGBColor(244, 248, 251)
LIGHT_BORDER = RGBColor(201, 212, 222)


def style_run(run, size=20, bold=False, color=TEXT_DARK, font_name="Segoe UI"):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_top_band(slide, title):
    band = slide.shapes.add_shape(
        1,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(0.78),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = PRIMARY
    band.line.fill.background()

    stripe = slide.shapes.add_shape(
        1,
        Inches(0),
        Inches(0.76),
        Inches(13.333),
        Inches(0.05),
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ACCENT
    stripe.line.fill.background()

    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.12), Inches(11.5), Inches(0.52)
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    style_run(run, size=25, bold=True, color=RGBColor(255, 255, 255))


def add_footer(slide, slide_no, total_slides):
    left = slide.shapes.add_textbox(Inches(0.6), Inches(7.1), Inches(8), Inches(0.3))
    left_tf = left.text_frame
    left_tf.clear()
    left_p = left_tf.paragraphs[0]
    left_run = left_p.add_run()
    left_run.text = "Harness + ServiceNow demo"
    style_run(left_run, size=11, color=TEXT_MID)

    right = slide.shapes.add_textbox(
        Inches(11.5), Inches(7.1), Inches(1.2), Inches(0.3)
    )
    right_tf = right.text_frame
    right_tf.clear()
    right_p = right_tf.paragraphs[0]
    right_p.alignment = PP_ALIGN.RIGHT
    right_run = right_p.add_run()
    right_run.text = f"{slide_no}/{total_slides}"
    style_run(right_run, size=11, color=TEXT_MID)


def add_bullet_slide(
    prs, title, bullets, slide_no, total_slides, screenshot_label=None
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_top_band(slide, title)

    content_box = slide.shapes.add_textbox(
        Inches(0.9), Inches(1.25), Inches(11.9), Inches(5.2)
    )
    tf = content_box.text_frame
    tf.clear()
    tf.word_wrap = True

    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(10)
        for run in p.runs:
            style_run(run, size=22, color=TEXT_DARK)

    if screenshot_label:
        shot = slide.shapes.add_shape(
            1,
            Inches(7.05),
            Inches(3.95),
            Inches(5.2),
            Inches(2.3),
        )
        shot.fill.solid()
        shot.fill.fore_color.rgb = LIGHT_FILL
        shot.line.color.rgb = LIGHT_BORDER
        shot.line.width = Pt(1.2)

        shot_text = shot.text_frame
        shot_text.clear()
        p = shot_text.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = screenshot_label
        style_run(run, size=14, color=TEXT_MID, bold=True)

        hint = shot_text.add_paragraph()
        hint.alignment = PP_ALIGN.CENTER
        hint_run = hint.add_run()
        hint_run.text = "Inserir screenshot da demo"
        style_run(hint_run, size=12, color=TEXT_MID)

    add_footer(slide, slide_no, total_slides)


def add_title_slide(prs, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY
    bg.line.fill.background()

    accent = slide.shapes.add_shape(
        1, Inches(0), Inches(5.7), Inches(13.333), Inches(1.8)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.2), Inches(11.8), Inches(2.2)
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Harness Platform"
    style_run(r, size=54, bold=True, color=RGBColor(255, 255, 255))

    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    r2.text = "Automacao inteligente para DevOps"
    style_run(r2, size=30, bold=False, color=RGBColor(227, 236, 244))

    p3 = tf.add_paragraph()
    p3.space_before = Pt(28)
    r3 = p3.add_run()
    r3.text = "Integracao com GitHub, AWS, Terraform e ServiceNow"
    style_run(r3, size=21, bold=True, color=RGBColor(255, 255, 255))

    date_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(6.0), Inches(7.4), Inches(0.5)
    )
    dtf = date_box.text_frame
    dtf.clear()
    dp = dtf.paragraphs[0]
    dr = dp.add_run()
    dr.text = f"Apresentacao para o time | {date.today().strftime('%d/%m/%Y')}"
    style_run(dr, size=14, bold=True, color=RGBColor(255, 255, 255))

    add_footer(slide, 1, total_slides)


prs = Presentation()
total = 8

add_title_slide(prs, total)

add_bullet_slide(
    prs,
    "Contexto",
    [
        "Harness automatiza CI/CD e operacoes de entrega de software.",
        "Complementa o GitHub: repositorio no GitHub, orquestracao no Harness.",
        "Objetivo: reduzir trabalho manual e elevar rastreabilidade em deploys.",
    ],
    slide_no=2,
    total_slides=total,
)

add_bullet_slide(
    prs,
    "Principais Recursos",
    [
        "CI/CD inteligente com pipelines reutilizaveis e aprovacoes controladas.",
        "Feature Flags para releases graduais e rollback rapido.",
        "Cloud Cost Management para visibilidade e otimizacao de custos.",
        "Seguranca e auditoria integradas para compliance do fluxo de entrega.",
    ],
    slide_no=3,
    total_slides=total,
)

add_bullet_slide(
    prs,
    "Ecossistema de Integracoes",
    [
        "Conectores nativos: GitHub, AWS, Terraform e fluxos em Python.",
        "ServiceNow out-of-the-box para ITSM e governanca de mudancas.",
        "Menos glue code, mais padronizacao entre squads.",
    ],
    slide_no=4,
    total_slides=total,
)

add_bullet_slide(
    prs,
    "Caso de Uso: Harness + ServiceNow",
    [
        "Pipeline aciona criacao de Change Request antes do deploy.",
        "Aprovacao no ServiceNow libera continuidade da execucao.",
        "Status final do deploy atualiza automaticamente o registro de mudanca.",
    ],
    slide_no=5,
    total_slides=total,
    screenshot_label="Screenshot: Change Request no ServiceNow",
)

add_bullet_slide(
    prs,
    "Demo Pipeline (4 minutos)",
    [
        "1) Mostrar connector Harness -> ServiceNow configurado.",
        "2) Executar pipeline com etapa de mudanca automatizada.",
        "3) Validar ticket e transicao de status no ServiceNow.",
        "4) Concluir com rastreabilidade: commit -> deploy -> change.",
    ],
    slide_no=6,
    total_slides=total,
    screenshot_label="Screenshot: execucao do pipeline no Harness",
)

add_bullet_slide(
    prs,
    "Beneficios para o Time",
    [
        "Menos operacao manual e menor risco em mudancas de producao.",
        "Compliance e auditoria nativos com trilha completa de execucao.",
        "Aderencia a stack atual: ServiceNow, AWS, Terraform, GitHub e Python.",
    ],
    slide_no=7,
    total_slides=total,
)

add_bullet_slide(
    prs,
    "Proximos Passos",
    [
        "Selecionar um piloto de deploy com Change Request automatizado.",
        "Medir lead time, taxa de sucesso e reducao de esforco manual.",
        "Escalar o padrao para outras squads apos validacao do piloto.",
        "Q&A",
    ],
    slide_no=8,
    total_slides=total,
)

output = "Harness_Platform_Apresentacao_v2.pptx"
prs.save(output)
print(output)
